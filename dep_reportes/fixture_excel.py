"""Lectura del Excel manual "Curva S - Proyecto <codigo>.xlsx" (valores cacheados de las formulas).

Entrega los insumos (tareas y horas) y los resultados que calculo el Excel, para validar metricas.py.
"""
from __future__ import annotations

import datetime as dt
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils import column_index_from_string

from dep_clickup.dates import ms_to_local, ms_to_local_date

from .metricas import Horas, PuntoSerie, TareaMetrica

FILAS_TAREAS = range(4, 200)       # 4..199, como los rangos SUM del Excel
FILAS_HORAS = range(3, 1001)       # 'Horas Cargadas'!A3:A1000
FILAS_RESUMEN = range(54, 76)      # Resumen!D54:G75


@dataclass(frozen=True)
class EntradaExcel:
    fila: int
    user_id: int
    start_ms: int
    duration_ms: int
    tarea_id: str
    tarea: str
    fecha_excel: dt.date | None     # columna A (calculada por el Excel desde el texto con offset)
    horas_excel: float              # columna B

    @property
    def fecha(self) -> dt.date:
        return ms_to_local_date(self.start_ms)  # type: ignore[return-value]

    @property
    def horas(self) -> float:
        return self.duration_ms / 3_600_000


@dataclass
class FixtureCurvaS:
    ruta: Path
    titulo: str
    programa: list[TareaMetrica]            # hoja Programado
    avance: list[TareaMetrica]              # hoja Avance
    entradas: list[EntradaExcel]            # hoja Horas Cargadas
    control: dt.date                        # Programado!K2
    fin: dt.date                            # Avance!M2
    inicio: dt.date                         # Programado!H2
    grilla: list[dt.date]                   # Resumen!D54:D75 (sin #N/A)
    serie: list[PuntoSerie]                 # Resumen!D:G tal como la calculo el Excel
    total_hh: float                         # Programado!M2
    total_hh_avance: float                  # Avance!T2
    avance_programado: float                # Programado!N2
    avance_real: float                      # Avance!U2
    hh_pendientes: float                    # Avance!D2
    fraccion_prog_tarea: dict[str, float] = field(default_factory=dict)   # Programado!N por tarea (ponderada)
    avance_real_tarea: dict[str, float] = field(default_factory=dict)     # Avance!U por tarea (ponderada)
    prog_diario_tarea: dict[str, dict[dt.date, float]] = field(default_factory=dict)  # Programado!P:DC
    pend_diario_tarea: dict[str, dict[dt.date, float]] = field(default_factory=dict)  # Avance!W:DJ
    prog_acumulado: dict[dt.date, float] = field(default_factory=dict)    # Programado fila 208
    pend_acumulado: dict[dt.date, float] = field(default_factory=dict)    # Avance fila 208
    padres: dict[str, str | None] = field(default_factory=dict)
    nombres: dict[str, str] = field(default_factory=dict)

    @property
    def horas(self) -> list[Horas]:
        return [Horas(e.fecha, e.horas, e.tarea_id) for e in self.entradas]


def _fecha(v: Any) -> dt.date | None:
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    return None


def _num(v: Any) -> float | None:
    if v is None or v == "" or (isinstance(v, str) and v.startswith("#")):
        return None
    if isinstance(v, str):
        v = v.strip().replace(",", ".")
        if not v:
            return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _str(v: Any) -> str:
    return "" if v is None else str(v).strip()


def _col(letra: str) -> int:
    return column_index_from_string(letra)


def _diario(ws, fila: int, col0: int, col1: int, fila_fechas: int) -> dict[dt.date, float]:
    out = {}
    for c in range(col0, col1 + 1):
        v = _num(ws.cell(fila, c).value)
        if v:
            out[_fecha(ws.cell(fila_fechas, c).value)] = v
    return out


def cargar(ruta: str | Path) -> FixtureCurvaS:
    ruta = Path(ruta)
    wb = openpyxl.load_workbook(ruta, data_only=True, read_only=False)
    P, A, H, R = wb["Programado"], wb["Avance"], wb["Horas Cargadas"], wb["Resumen"]

    programa, fracc, prog_diario, padres, nombres = [], {}, {}, {}, {}
    p0, p1 = _col("P"), _col("DC")
    for r in FILAS_TAREAS:
        tid = _str(P.cell(r, 1).value)
        if not tid:
            continue
        t = TareaMetrica(
            id=tid, nombre=_str(P.cell(r, 2).value), parent=_str(P.cell(r, 3).value) or None,
            start=_fecha(P.cell(r, _col("H")).value), due=_fecha(P.cell(r, _col("I")).value),
            hh=_num(P.cell(r, _col("M")).value), avance=_num(P.cell(r, _col("L")).value),
        )
        programa.append(t)
        padres[tid], nombres[tid] = t.parent, t.nombre
        n = _num(P.cell(r, _col("N")).value)
        if n is not None and r >= 5:  # N2 = SUM(N5:N199)
            fracc[tid] = n
        prog_diario[tid] = _diario(P, r, p0, p1, 3)

    avance, u_tarea, pend_diario = [], {}, {}
    w0, w1 = _col("W"), _col("DJ")
    for r in FILAS_TAREAS:
        tid = _str(A.cell(r, _col("F")).value)
        if not tid:
            continue
        t = TareaMetrica(
            id=tid, nombre=_str(A.cell(r, _col("G")).value), parent=_str(A.cell(r, _col("H")).value) or None,
            start=_fecha(A.cell(r, _col("L")).value), due=_fecha(A.cell(r, _col("M")).value),
            hh=_num(A.cell(r, _col("T")).value), avance=_num(A.cell(r, _col("S")).value),
        )
        avance.append(t)
        padres.setdefault(tid, t.parent)
        nombres.setdefault(tid, t.nombre)
        u = _num(A.cell(r, _col("U")).value)
        if u is not None and r >= 5:  # U2 = SUM(U5:U199)
            u_tarea[tid] = u
        pend_diario[tid] = _diario(A, r, w0, w1, 3)

    entradas = []
    for r in FILAS_HORAS:
        start = H.cell(r, _col("J")).value
        if start in (None, ""):
            continue
        entradas.append(EntradaExcel(
            fila=r, user_id=int(H.cell(r, _col("D")).value), start_ms=int(start),
            duration_ms=int(H.cell(r, _col("N")).value), tarea_id=_str(H.cell(r, _col("V")).value),
            tarea=_str(H.cell(r, _col("W")).value), fecha_excel=_fecha(H.cell(r, 1).value),
            horas_excel=_num(H.cell(r, 2).value) or 0.0,
        ))

    grilla, serie = [], []
    for r in FILAS_RESUMEN:
        d = _fecha(R.cell(r, _col("D")).value)
        if d is None:
            continue
        grilla.append(d)
        serie.append(PuntoSerie(d, _num(R.cell(r, _col("E")).value) or 0.0,
                                _num(R.cell(r, _col("F")).value), _num(R.cell(r, _col("G")).value)))

    def acumulado(ws, c0, c1):
        return {_fecha(ws.cell(207, c).value): _num(ws.cell(208, c).value) or 0.0 for c in range(c0, c1 + 1)}

    return FixtureCurvaS(
        ruta=ruta, titulo=_str(P["P1"].value),
        programa=programa, avance=avance, entradas=entradas,
        control=_fecha(P["K2"].value), fin=_fecha(A["M2"].value), inicio=_fecha(P["H2"].value),
        grilla=grilla, serie=serie,
        total_hh=_num(P["M2"].value), total_hh_avance=_num(A["T2"].value),
        avance_programado=_num(P["N2"].value), avance_real=_num(A["U2"].value),
        hh_pendientes=_num(A["D2"].value),
        fraccion_prog_tarea=fracc, avance_real_tarea=u_tarea,
        prog_diario_tarea=prog_diario, pend_diario_tarea=pend_diario,
        prog_acumulado=acumulado(P, p0, p1), pend_acumulado=acumulado(A, w0, w1),
        padres=padres, nombres=nombres,
    )


def fecha_export(ruta: str | Path) -> dt.datetime:
    """Momento aproximado del export: fecha de modificacion del archivo."""
    return dt.datetime.fromtimestamp(Path(ruta).stat().st_mtime).astimezone(ms_to_local(0).tzinfo)
